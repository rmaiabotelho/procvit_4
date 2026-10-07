import os
import re
import shutil
import sqlite3
import sys
import time
import subprocess
from datetime import datetime
import tkinter as tk
from tkinter import filedialog
from typing import Dict, List, Optional, Tuple

from colorama import Fore, Style, init
from docx import Document

# Inicializa o Colorama
init(autoreset=True)

RESET = Style.RESET_ALL
VERMELHO = Fore.RED + Style.BRIGHT
VERDE = Fore.GREEN + Style.BRIGHT
AMARELO = Fore.YELLOW + Style.BRIGHT
CIANO = Fore.CYAN + Style.BRIGHT
NEGRITO = Style.BRIGHT
AZUL = Fore.BLUE + Style.BRIGHT
MAGENTA = Fore.MAGENTA + Style.BRIGHT
CINZA_ESCURO = Fore.BLACK + Style.BRIGHT


def limpar_terminal() -> None:
    comando = "cls" if os.name == "nt" else "clear"
    subprocess.run(comando, shell=True)


def menu_pausa() -> None:
    input(f"\n{AMARELO}Pressione Enter para voltar ao menu...{RESET}")


def input_cancelavel(prompt: str) -> Tuple[str, bool]:
    entrada = input(f"{prompt} | [Q] Cancelar: ").strip()
    if entrada.lower() == "q":
        return "", True
    return entrada, False


def hoje() -> str:
    return datetime.now().strftime("%Y-%m-%d")


def formatar_data(data_str: str) -> str:
    numeros = "".join(filter(str.isdigit, data_str))
    if len(numeros) != 8:
        return ""

    for fmt in ("%Y%m%d", "%d%m%Y"):
        try:
            dt = datetime.strptime(numeros, fmt)
            return dt.strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


def validar_formatar_doc(doc: str) -> str:
    n = "".join(filter(str.isdigit, doc))
    if len(n) not in (11, 14) or len(set(n)) == 1:
        return ""

    def calc_digito(s: str, pesos: List[int]) -> int:
        soma = sum(int(a) * b for a, b in zip(s, pesos))
        resto = soma % 11
        return 0 if resto < 2 else 11 - resto

    if len(n) == 11:
        d1 = calc_digito(n[:9], list(range(10, 1, -1)))
        d2 = calc_digito(n[:9] + str(d1), list(range(11, 1, -1)))
        if n[9:] != f"{d1}{d2}":
            return ""
        return f"{n[:3]}.{n[3:6]}.{n[6:9]}-{n[9:]}"
    else:
        pesos1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
        pesos2 = [6] + pesos1
        d1 = calc_digito(n[:12], pesos1)
        d2 = calc_digito(n[:12] + str(d1), pesos2)
        if n[12:] != f"{d1}{d2}":
            return ""
        return f"{n[:2]}.{n[2:5]}.{n[5:8]}/{n[8:12]}-{n[12:]}"


def formatar_telefone(telefone: str) -> str:
    fone = "".join(filter(str.isdigit, telefone))
    if not fone.startswith("55"):
        fone = "55" + fone
    return f"+{fone}" if len(fone) >= 12 else ""


def formatar_endereco(endereco: str) -> str:
    minusculas = {"de", "do", "da", "dos", "das", "e", "com", "em", "apto", "apt", "sala", "bloco", "quadra", "lote", "lt.", "qd.", "lt", "qd"}
    maiusculas = {"rj", "sp", "mg", "es", "pr", "sc", "rs", "ms", "mt", "go", "df", "to", "pa", "am", "ro", "ac", "ap", "rr", "ma", "pi", "ce", "rn", "pb", "pe", "al", "se", "ba", "cep", "bpc", "loas"}
    
    palavras = endereco.lower().split()
    resultado = []
    for p in palavras:
        if p in maiusculas:
            resultado.append(p.upper())
        elif p in minusculas:
            resultado.append(p)
        else:
            resultado.append(p.capitalize())
    return " ".join(resultado)


def formatar_processo(processo: str) -> str:
    n = "".join(filter(str.isdigit, processo))
    n = n[-20:].zfill(20)
    return f"{n[0:7]}-{n[7:9]}.{n[9:13]}.{n[13]}.{n[14:16]}.{n[16:20]}"


def selecionar_arquivo(titulo: str = "Selecione o arquivo", extensoes: List[Tuple[str, str]] = [("Documentos Word", "*.docx")]) -> str:
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    caminho = filedialog.askopenfilename(title=titulo, filetypes=extensoes)
    root.destroy()
    return caminho


def selecionar_salvar(nome_padrao: str = "arquivo.docx", extensoes: List[Tuple[str, str]] = [("Documentos Word", "*.docx")]) -> str:
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    caminho = filedialog.asksaveasfilename(title="Escolha onde salvar", initialfile=nome_padrao, filetypes=extensoes)
    root.destroy()
    return caminho


def criar_db(caminho: str) -> None:
    try:
        with sqlite3.connect(caminho) as conn:
            cursor = conn.cursor()
            cursor.execute("PRAGMA foreign_keys = ON;")
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS clientes (
                    NOME TEXT NOT NULL,
                    DATA_NASCIMENTO TEXT,
                    IDENTIDADE TEXT,
                    CPF_CNPJ TEXT PRIMARY KEY,
                    ESTADO_CIVIL TEXT,
                    PROFISSAO TEXT,
                    ENDERECO TEXT,
                    TELEFONE TEXT,
                    EMAIL TEXT,
                    SENHA_GOV TEXT,
                    OBS TEXT
                )
            """)
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS processos (
                    PROCESSO TEXT PRIMARY KEY,
                    CPF_CNPJ TEXT NOT NULL,
                    CLIENTE TEXT,
                    PARTE_CONTRARIA TEXT,
                    CARTORIO TEXT,
                    SITUACAO TEXT,
                    ULTIMA_VERIFICACAO TEXT,
                    DISTRIBUICAO TEXT,
                    OBS TEXT,
                    FOREIGN KEY (CPF_CNPJ) REFERENCES clientes (CPF_CNPJ)
                )
            """)
    except Exception as e:
        print(f"{VERMELHO}[ERRO] Falha ao criar tabelas: {e}{RESET}")


def conectar_db(caminho: str) -> sqlite3.Connection:
    conn = sqlite3.connect(caminho)
    conn.row_factory = sqlite3.Row
    return conn


def backup_db(origem: str, destino: str) -> None:
    if not os.path.exists(origem):
        print(f"{VERMELHO}[Erro] Arquivo de origem não encontrado.{RESET}")
        return
    try:
        shutil.copy2(origem, destino)
        print(f"{VERDE}BACKUP CONCLUÍDO!\nLocal: {destino}{RESET}")
    except Exception as e:
        print(f"{VERMELHO}[Erro] Falha ao realizar backup: {e}{RESET}")


def substituir_texto_docx(doc: Document, substituicoes: Dict[str, str]) -> None:
    for p in doc.paragraphs:
        for ch, vl in substituicoes.items():
            if ch in p.text:
                p.text = p.text.replace(ch, vl)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for ch, vl in substituicoes.items():
                        if ch in p.text:
                            p.text = p.text.replace(ch, vl)


def preparar_dados_documento(conn: sqlite3.Connection, busca: str) -> Dict[str, str]:
    dados = {}
    cursor = conn.cursor()
    
    cursor.execute("""
        SELECT PROCESSO, CPF_CNPJ, CLIENTE, PARTE_CONTRARIA, CARTORIO 
        FROM processos WHERE PROCESSO = ? OR CPF_CNPJ = ? LIMIT 1
    """, (busca, busca))
    row = cursor.fetchone()

    cpf_cnpj_alvo = None

    if row:
        cpf_cnpj_alvo = row["CPF_CNPJ"]
        dados["<<PROCESSO>>"] = row["PROCESSO"] or ""
        dados["<<CPF_CNPJ>>"] = row["CPF_CNPJ"] or ""
        dados["<<CLIENTE>>"] = (row["CLIENTE"] or "").upper()
        dados["<<PARTECONTRARIA>>"] = (row["PARTE_CONTRARIA"] or "").upper()
        dados["<<CARTÓRIO>>"] = (row["CARTORIO"] or "").upper()
    else:
        cursor.execute("SELECT CPF_CNPJ, NOME FROM clientes WHERE CPF_CNPJ = ?", (busca,))
        row_cli = cursor.fetchone()
        if row_cli:
            cpf_cnpj_alvo = row_cli["CPF_CNPJ"]
            dados["<<CPF_CNPJ>>"] = row_cli["CPF_CNPJ"] or ""
            dados["<<CLIENTE>>"] = (row_cli["NOME"] or "").upper()
            dados["<<PROCESSO>>"] = ""
            dados["<<PARTECONTRARIA>>"] = ""
            dados["<<CARTÓRIO>>"] = ""

    if cpf_cnpj_alvo:
        dt = datetime.now()
        meses = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", 
                 "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
        
        dados["<<DD>>"] = dt.strftime("%d")
        dados["<<MM>>"] = meses[dt.month - 1]
        dados["<<AAAA>>"] = dt.strftime("%Y")

        cursor.execute("""
            SELECT ENDERECO, PROFISSAO, ESTADO_CIVIL, IDENTIDADE, EMAIL, DATA_NASCIMENTO 
            FROM clientes WHERE CPF_CNPJ = ?
        """, (cpf_cnpj_alvo,))
        row_c = cursor.fetchone()
        
        if row_c:
            dados["<<ENDEREÇO>>"] = row_c["ENDERECO"] or ""
            dados["<<PROFISSÃO>>"] = (row_c["PROFISSAO"] or "").lower()
            dados["<<ESTADO CIVIL>>"] = (row_c["ESTADO_CIVIL"] or "").lower()
            dados["<<RG>>"] = row_c["IDENTIDADE"] or ""
            dados["<<EMAIL>>"] = (row_c["EMAIL"] or "").lower()
            
            data_nasc_raw = row_c["DATA_NASCIMENTO"] or ""
            if len(data_nasc_raw) == 10 and "-" in data_nasc_raw:
                partes = data_nasc_raw.split("-")
                dados["<<DATA_NASCIMENTO>>"] = f"{partes[2]}/{partes[1]}/{partes[0]}"
                dados["<<NASCIMENTO>>"] = f"{partes[2]}/{partes[1]}/{partes[0]}"
            else:
                dados["<<DATA_NASCIMENTO>>"] = data_nasc_raw
                dados["<<NASCIMENTO>>"] = data_nasc_raw

    return dados


def preencher_docx(conn: sqlite3.Connection) -> None:
    print(f"{AZUL}=========== GERADOR DE DOCUMENTOS PROCVIT ===========")
    busca, cancelou = input_cancelavel(f"{AZUL}Digite o Processo ou CPF/CNPJ{RESET}")
    if cancelou or not busca:
        return

    dados = preparar_dados_documento(conn, busca)
    if not dados:
        print(f"{VERMELHO}[!] Nenhum registro encontrado.{RESET}")
        time.sleep(1.5)
        return

    print(f"{AMARELO}\n> Escolha o modelo .docx...{RESET}")
    modelo = selecionar_arquivo("Selecione o Modelo de Documento")
    if not modelo:
        return

    sugestao = f"Gerado_{dados.get('<<CLIENTE>>', 'doc')}.docx"
    print(f"{AMARELO}> Onde deseja salvar o documento preenchido?{RESET}")
    destino = selecionar_salvar(sugestao)

    if destino:
        try:
            doc = Document(modelo)
            substituir_texto_docx(doc, dados)
            doc.save(destino)
            print(f"\n{VERDE}[OK] Documento gerado com sucesso!{RESET}")
            
            if os.name == 'nt':
                os.startfile(destino)
            else:
                subprocess.run(["xdg-open", destino], check=False)

        except Exception as e:
            print(f"{VERMELHO}[!] Erro ao gerar/abrir documento: {e}{RESET}")
        time.sleep(2)


def cadastrar_cliente(conn: sqlite3.Connection) -> None:
    while True:
        limpar_terminal()
        print("\n----------- NOVO CADASTRO DE CLIENTE -----------")
        doc_raw, cancel = input_cancelavel("🪪 CPF/CNPJ")
        if cancel: 
            return
        
        cpf = validar_formatar_doc(doc_raw)
        if not cpf:
            print(f"{VERMELHO}[!] Documento inválido.{RESET}")
            time.sleep(1.5)
            continue

        cursor = conn.cursor()
        cursor.execute("SELECT 1 FROM clientes WHERE CPF_CNPJ = ?", (cpf,))
        if cursor.fetchone():
            print(f"{AMARELO}[!] Cliente já cadastrado.{RESET}")
            time.sleep(1.5)
            continue

        nome, c = input_cancelavel("👤 Nome Completo") 
        if c: return

        while True:
            nasc_raw, c = input_cancelavel("📅 Data Nasc (DDMMAAAA)") 
            if c: return
            data_nasc = formatar_data(nasc_raw)
            if data_nasc:
                break
            print(f"{VERMELHO}[!] Data inválida. Use o formato DDMMAAAA ou DD/MM/AAAA.{RESET}")
            time.sleep(1)

        rg, c = input_cancelavel("🪪 RG") 
        if c: return
        end, c = input_cancelavel("📍 Endereço") 
        if c: return
        tel, c = input_cancelavel("📞 Telefone") 
        if c: return
        email, c = input_cancelavel("📧 Email") 
        if c: return
        est_civil, c = input_cancelavel("📋 Estado Civil") 
        if c: return
        prof, c = input_cancelavel("💼 Profissão") 
        if c: return
        senha, c = input_cancelavel("🔑 Senha GOV") 
        if c: return
        obs, c = input_cancelavel("📝 Observações") 
        if c: return

        try:
            # CORREÇÃO: Ordem explícita correspondendo à tupla de parâmetros
            cursor.execute("""
                INSERT INTO clientes (
                    CPF_CNPJ, NOME, DATA_NASCIMENTO, IDENTIDADE, 
                    ESTADO_CIVIL, PROFISSAO, ENDERECO, TELEFONE, 
                    EMAIL, SENHA_GOV, OBS
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                cpf, nome.upper().strip(), data_nasc, rg.upper().strip(), 
                est_civil.upper().strip(), prof.upper().strip(), formatar_endereco(end), 
                formatar_telefone(tel), email.lower().strip(), senha.strip(), obs.strip()
            ))
            conn.commit()
            print(f"\n{VERDE}[OK] Cliente cadastrado com sucesso!{RESET}")
        except Exception as e:
            print(f"{VERMELHO}[ERRO]: {e}{RESET}")
        
        time.sleep(2)
        cont, cancelou = input_cancelavel("\nCadastrar outro? [Enter] Sim")
        if cancelou or cont.lower() == "q": 
            break


def buscar_cliente(conn: sqlite3.Connection) -> None:
    while True:
        limpar_terminal()
        termo_input = input("\n 👤 Buscar Cliente (Enter para todos) | [Q] Voltar ao Menu: ").strip()
        if termo_input.lower() == "q": 
            return

        # Se pressionar Enter sem digitar nada, busca todos
        termo = f"%{termo_input}%" if termo_input else "%"
        pagina = 0
        limite = 5

        while True:
            limpar_terminal()
            cursor = conn.cursor()
            cursor.execute("""
                SELECT COUNT(*) FROM clientes 
                WHERE NOME LIKE ? OR CPF_CNPJ LIKE ? OR PROFISSAO LIKE ? OR EMAIL LIKE ?
            """, (termo, termo, termo, termo))
            total = cursor.fetchone()[0]

            if total == 0:
                print(f"\n{AMARELO}[!] Nenhum registro encontrado.{RESET}")
                time.sleep(1.5)
                break

            total_paginas = max(1, (total + limite - 1) // limite)

            cursor.execute("""
                SELECT * FROM clientes 
                WHERE NOME LIKE ? OR CPF_CNPJ LIKE ? OR PROFISSAO LIKE ? OR EMAIL LIKE ?
                ORDER BY NOME ASC LIMIT ? OFFSET ?
            """, (termo, termo, termo, termo, limite, pagina * limite))
            rows = cursor.fetchall()

            print(f"\n{AZUL}===================== CLIENTES | Pág: {pagina + 1}/{total_paginas} | TOTAL: {total} ====================={RESET}")
            
            for i, r in enumerate(rows):
                print("-" * 80)
                rg = r['IDENTIDADE'] if r['IDENTIDADE'] else 'N/A'
                profissao = r['PROFISSAO'] if r['PROFISSAO'] else 'N/A'
                est_civil = r['ESTADO_CIVIL'] if r['ESTADO_CIVIL'] else 'N/A'
                tel = r['TELEFONE'] if r['TELEFONE'] else 'N/A'
                email = r['EMAIL'] if r['EMAIL'] else 'N/A'
                end = r['ENDERECO'] if r['ENDERECO'] else 'N/A'
                senha_gov = r['SENHA_GOV'] if r['SENHA_GOV'] else 'N/A'
                obs = r['OBS'] if r['OBS'] else 'N/A'
                
                data_nasc_raw = r['DATA_NASCIMENTO'] or ''
                if len(data_nasc_raw) == 10 and "-" in data_nasc_raw:
                    p = data_nasc_raw.split("-")
                    dt_nasc = f"{p[2]}/{p[1]}/{p[0]}"
                else:
                    dt_nasc = data_nasc_raw or 'N/A'

                print(f"{VERDE}{i + 1}.{RESET} 👤 {AMARELO}{r['NOME']}{RESET}")
                print(f"   🪪  CPF/CNPJ: {r['CPF_CNPJ']} | RG: {rg} | 📅 NASC: {dt_nasc}")
                print(f"   💼 PROFISSÃO: {profissao} | 💍 EST. CIVIL: {est_civil}")
                print(f"   📞 TEL: {tel} | ✉️  EMAIL: {email}")
                print(f"   🏠 ENDEREÇO: {end}")
                print(f"   🔑 SENHA GOV: {senha_gov}")
                print(f"   📝 OBS: {obs}")
            
            print("-" * 80)
            print("[<] Ant | [>] Próx | [E+Nº] Editar | [G] Gerar Doc | [Q] Nova Busca")
            acao = input("Comando: ").strip().lower()

            if acao == "q": 
                break
            elif acao in (">", ".") and pagina + 1 < total_paginas: 
                pagina += 1
            elif acao in ("<", ",") and pagina > 0: 
                pagina -= 1
            elif acao == "g": 
                preencher_docx(conn)
            elif acao.startswith("e"):
                try:
                    idx = int(acao.replace("e", "").strip()) - 1
                    if 0 <= idx < len(rows):
                        editar_cliente(conn, rows[idx]['CPF_CNPJ'])
                except ValueError: 
                    pass


def editar_cliente(conn: sqlite3.Connection, cpf: str = "") -> None:
    cliente_cpf = cpf

    if not cliente_cpf:
        busca, cancelou = input_cancelavel("Digite o CPF/CNPJ do cliente para editar")
        if cancelou or not busca:
            return
        cliente_cpf = validar_formatar_doc(busca)

    cursor = conn.cursor()
    cursor.execute("SELECT * FROM clientes WHERE CPF_CNPJ = ?", (cliente_cpf,))
    row = cursor.fetchone()

    if not row:
        print(f"{VERMELHO}[!] Cliente não encontrado.{RESET}")
        time.sleep(1.5)
        return

    dados = dict(row)
    houve_alteracao = False

    while True:
        limpar_terminal()
        print(f"\n{AZUL}================ EDITAR CLIENTE ================{RESET}")
        print(f"CPF/CNPJ (Chave)     : {dados['CPF_CNPJ']}")
        print(f"1. Nome              : {dados['NOME'] or 'N/A'}")
        print(f"2. Data Nascimento   : {dados['DATA_NASCIMENTO'] or 'N/A'}")
        print(f"3. RG / Identidade   : {dados['IDENTIDADE'] or 'N/A'}")
        print(f"4. Estado Civil      : {dados['ESTADO_CIVIL'] or 'N/A'}")
        print(f"5. Profissão         : {dados['PROFISSAO'] or 'N/A'}")
        print(f"6. Endereço          : {dados['ENDERECO'] or 'N/A'}")
        print(f"7. Telefone          : {dados['TELEFONE'] or 'N/A'}")
        print(f"8. Email             : {dados['EMAIL'] or 'N/A'}")
        print(f"9. Senha GOV         : {dados['SENHA_GOV'] or 'N/A'}")
        print(f"10. Observações      : {dados['OBS'] or 'N/A'}")
        print(f"\n0. {VERDE}SALVAR ALTERAÇÕES{RESET}")
        print("q. Cancelar e Sair")

        opcao, cancelou = input_cancelavel("\nEscolha o campo para editar")
        if cancelou or opcao.lower() == "q":
            print("\nEdição cancelada.")
            time.sleep(1)
            return
        if opcao == "0":
            break

        if opcao == "1":
            v, _ = input_cancelavel(f"Novo Nome [{dados['NOME'] or ''}]")
            if v: 
                dados['NOME'] = v.upper().strip()
                houve_alteracao = True
        elif opcao == "2":
            v, _ = input_cancelavel(f"Nova Data Nasc (DDMMAAAA) [{dados['DATA_NASCIMENTO'] or ''}]")
            if v:
                dt = formatar_data(v)
                if dt:
                    dados['DATA_NASCIMENTO'] = dt
                    houve_alteracao = True
                else:
                    print(f"{VERMELHO}[!] Data inválida.{RESET}")
                    time.sleep(1)
        elif opcao == "3":
            v, _ = input_cancelavel(f"Novo RG [{dados['IDENTIDADE'] or ''}]")
            if v: 
                dados['IDENTIDADE'] = v.upper().strip()
                houve_alteracao = True
        elif opcao == "4":
            v, _ = input_cancelavel(f"Novo Estado Civil [{dados['ESTADO_CIVIL'] or ''}]")
            if v: 
                dados['ESTADO_CIVIL'] = v.upper().strip()
                houve_alteracao = True
        elif opcao == "5":
            v, _ = input_cancelavel(f"Nova Profissão [{dados['PROFISSAO'] or ''}]")
            if v: 
                dados['PROFISSAO'] = v.upper().strip()
                houve_alteracao = True
        elif opcao == "6":
            v, _ = input_cancelavel(f"Novo Endereço [{dados['ENDERECO'] or ''}]")
            if v: 
                dados['ENDERECO'] = formatar_endereco(v)
                houve_alteracao = True
        elif opcao == "7":
            v, _ = input_cancelavel(f"Novo Telefone [{dados['TELEFONE'] or ''}]")
            if v: 
                dados['TELEFONE'] = formatar_telefone(v)
                houve_alteracao = True
        elif opcao == "8":
            v, _ = input_cancelavel(f"Novo Email [{dados['EMAIL'] or ''}]")
            if v: 
                dados['EMAIL'] = v.lower().strip()
                houve_alteracao = True
        elif opcao == "9":
            v, _ = input_cancelavel(f"Nova Senha GOV [{dados['SENHA_GOV'] or ''}]")
            if v: 
                dados['SENHA_GOV'] = v.strip()
                houve_alteracao = True
        elif opcao == "10":
            v, _ = input_cancelavel(f"Nova Observação [{dados['OBS'] or ''}]")
            if v: 
                dados['OBS'] = v.strip()
                houve_alteracao = True
        else:
            print(f"{AMARELO}Opção inválida.{RESET}")
            time.sleep(1)

    if houve_alteracao:
        try:
            cursor.execute("""
                UPDATE clientes SET 
                    NOME = ?, DATA_NASCIMENTO = ?, IDENTIDADE = ?, 
                    ESTADO_CIVIL = ?, PROFISSAO = ?, ENDERECO = ?, 
                    TELEFONE = ?, EMAIL = ?, SENHA_GOV = ?, OBS = ?
                WHERE CPF_CNPJ = ?
            """, (
                dados['NOME'], dados['DATA_NASCIMENTO'], dados['IDENTIDADE'],
                dados['ESTADO_CIVIL'], dados['PROFISSAO'], dados['ENDERECO'],
                dados['TELEFONE'], dados['EMAIL'], dados['SENHA_GOV'],
                dados['OBS'], cliente_cpf
            ))
            conn.commit()
            print(f"\n{VERDE}[OK] Cliente atualizado com sucesso!{RESET}")
        except Exception as e:
            print(f"{VERMELHO}[ERRO] Falha ao atualizar: {e}{RESET}")
    else:
        print("\nNenhuma alteração foi feita.")
    
    time.sleep(1.5)


def cadastrar_processo(conn: sqlite3.Connection) -> None:
    while True:
        limpar_terminal()
        print("------- NOVO CADASTRO DE PROCESSO -------")
        proc_raw, c = input_cancelavel("⚖️ Número do Processo")
        if c: 
            return
        num_proc = formatar_processo(proc_raw)

        cursor = conn.cursor()
        cursor.execute("SELECT 1 FROM processos WHERE PROCESSO = ?", (num_proc,))
        if cursor.fetchone():
            print(f"{AMARELO}[!] Processo já cadastrado.{RESET}")
            time.sleep(1.5)
            continue

        doc_cli, c = input_cancelavel("🪪 CPF/CNPJ do Cliente")
        if c: 
            return
            
        cpf_cli = validar_formatar_doc(doc_cli)
        if not cpf_cli:
            print(f"{VERMELHO}[!] CPF/CNPJ inválido.{RESET}")
            time.sleep(1.5)
            continue
        
        cursor.execute("SELECT NOME FROM clientes WHERE CPF_CNPJ = ?", (cpf_cli,))
        cli_row = cursor.fetchone()
        if not cli_row:
            print(f"{VERMELHO}[!] Cliente não cadastrado no sistema.{RESET}")
            time.sleep(2)
            return

        parte, c = input_cancelavel("⚔️ Parte Contrária")
        if c: 
            return
        cartorio, c = input_cancelavel("🏛️ Cartório/Vara")
        if c: 
            return
        
        dt_dist, c = input_cancelavel("📅 Data Distribuição (DDMMAAAA - Deixe em branco p/ Hoje)")
        if c: return
        data_distribuicao = formatar_data(dt_dist) if dt_dist else hoje()

        obs, c = input_cancelavel("📝 Observações")
        if c: 
            return

        try:
            cursor.execute("""
                INSERT INTO processos (
                    PROCESSO, CPF_CNPJ, CLIENTE, PARTE_CONTRARIA, 
                    CARTORIO, SITUACAO, ULTIMA_VERIFICACAO, DISTRIBUICAO, OBS
                ) VALUES (?, ?, ?, ?, ?, 'ATIVO', ?, ?, ?)
            """, (
                num_proc, cpf_cli, cli_row['NOME'], parte.upper().strip(), 
                cartorio.upper().strip(), hoje(), data_distribuicao, obs.upper().strip()
            ))
            conn.commit()
            print(f"\n{VERDE}[OK] Processo cadastrado com sucesso!{RESET}")
        except Exception as e:
            print(f"{VERMELHO}[ERRO]: {e}{RESET}")
            
        time.sleep(2)
        break


def buscar_processo(conn: sqlite3.Connection) -> None:
    while True:
        limpar_terminal()
        termo_input = input("\n 📄 Buscar Processo (Enter para todos) | [Q] Voltar ao Menu: ").strip()
        if termo_input.lower() == "q": 
            return

        # Se pressionar Enter sem digitar nada, busca todos
        termo = f"%{termo_input}%" if termo_input else "%"
        pagina = 0
        limite = 5

        while True:
            limpar_terminal()
            cursor = conn.cursor()
            cursor.execute("""
                SELECT COUNT(*) FROM processos 
                WHERE PROCESSO LIKE ? OR CLIENTE LIKE ? OR PARTE_CONTRARIA LIKE ? OR CPF_CNPJ LIKE ?
            """, (termo, termo, termo, termo))
            total = cursor.fetchone()[0]
            
            if total == 0:
                print(f"\n{AMARELO}[!] Nenhum processo encontrado.{RESET}")
                time.sleep(1.5)
                break

            total_paginas = max(1, (total + limite - 1) // limite)

            cursor.execute("""
                SELECT * FROM processos 
                WHERE PROCESSO LIKE ? OR CLIENTE LIKE ? OR PARTE_CONTRARIA LIKE ? OR CPF_CNPJ LIKE ?
                ORDER BY DISTRIBUICAO DESC LIMIT ? OFFSET ?
            """, (termo, termo, termo, termo, limite, pagina * limite))
            rows = cursor.fetchall()

            print(f"\n{AZUL}===================== PROCESSOS | Pág: {pagina + 1}/{total_paginas} | TOTAL: {total} ====================={RESET}")
            for i, r in enumerate(rows):
                cartorio = r['CARTORIO'] if r['CARTORIO'] else 'N/A'
                distribuicao = r['DISTRIBUICAO'] if r['DISTRIBUICAO'] else 'N/A'
                ult_verif = r['ULTIMA_VERIFICACAO'] if r['ULTIMA_VERIFICACAO'] else 'N/A'
                cpf_cnpj = r['CPF_CNPJ'] if r['CPF_CNPJ'] else 'N/A'
                contraria = r['PARTE_CONTRARIA'] if r['PARTE_CONTRARIA'] else 'N/A'
                obs = r['OBS'] if r['OBS'] else 'N/A'

                sit_str = r['SITUACAO'] if r['SITUACAO'] else 'N/A'
                if sit_str == "ATIVO":
                    sit_cor = f"{VERDE}{sit_str}{RESET}"
                elif sit_str == "CONCLUIDO":
                    sit_cor = f"{VERMELHO}{sit_str}{RESET}"
                else:
                    sit_cor = sit_str

                print("-" * 80)
                print(f"{VERDE}{i + 1}.{RESET} 📄 PROC: {AMARELO}{r['PROCESSO']}{RESET} | SITUAÇÃO: {sit_cor}")
                print(f"   👤 CLIENTE: {r['CLIENTE']} | CPF/CNPJ: {cpf_cnpj}")
                print(f"   ⚔️  PARTE CONTRÁRIA: {contraria}")
                print(f"   🏛️  CARTÓRIO/VARA: {cartorio}")
                print(f"   📅 DISTRIBUIÇÃO: {distribuicao} | 🔍 ÚLT. VERIF: {ult_verif}")
                print(f"   📝 OBS: {obs}")

            print("-" * 80)
            print("[<] Ant | [>] Próx | [A+Nº] Visto | [E+Nº] Editar | [G] Gerar Doc | [Q] Nova Busca")
            acao = input("Comando: ").strip().lower()

            if acao == "q": 
                break
            elif acao in (">", ".") and pagina + 1 < total_paginas: 
                pagina += 1
            elif acao in ("<", ",") and pagina > 0: 
                pagina -= 1
            elif acao == "g": 
                preencher_docx(conn)
            elif acao.startswith("a"):
                try:
                    idx = int(acao.replace("a", "").strip()) - 1
                    if 0 <= idx < len(rows):
                        cursor.execute("UPDATE processos SET ULTIMA_VERIFICACAO = ? WHERE PROCESSO = ?", (hoje(), rows[idx]['PROCESSO']))
                        conn.commit()
                        print(f"{VERDE}[OK] Data de verificação atualizada!{RESET}")
                        time.sleep(0.8)
                except ValueError: 
                    pass
            elif acao.startswith("e"):
                try:
                    idx = int(acao.replace("e", "").strip()) - 1
                    if 0 <= idx < len(rows):
                        editar_processo(conn, rows[idx]['PROCESSO'])
                except ValueError: 
                    pass


def editar_processo(conn: sqlite3.Connection, proc_direto: str = "") -> None:
    processo_id = proc_direto

    if not processo_id:
        busca, cancelou = input_cancelavel("Digite o número do processo para editar")
        if cancelou or not busca:
            return
        processo_id = formatar_processo(busca)

    cursor = conn.cursor()
    cursor.execute("""
        SELECT CLIENTE, PARTE_CONTRARIA, CARTORIO, SITUACAO, DISTRIBUICAO, OBS 
        FROM processos WHERE PROCESSO = ?
    """, (processo_id,))
    row = cursor.fetchone()

    if not row:
        print(f"{VERMELHO}[!] Processo não encontrado.{RESET}")
        time.sleep(1.5)
        return

    dados = dict(row)
    houve_alteracao = False

    while True:
        limpar_terminal()
        print(f"\n{AZUL}================ EDITAR PROCESSO ================{RESET}")
        print(f"Processo (Chave)     : {processo_id}")
        print(f"1. Cliente (Nome)    : {dados['CLIENTE'] or 'N/A'}")
        print(f"2. Parte Contrária   : {dados['PARTE_CONTRARIA'] or 'N/A'}")
        print(f"3. Cartório / Vara   : {dados['CARTORIO'] or 'N/A'}")
        print(f"4. Situação          : {dados['SITUACAO'] or 'N/A'}")
        print(f"5. Distribuição      : {dados['DISTRIBUICAO'] or 'N/A'}")
        print(f"6. Observações       : {dados['OBS'] or 'N/A'}")
        print(f"\n0. {VERDE}SALVAR ALTERAÇÕES{RESET}")
        print("q. Cancelar e Sair")

        opcao, cancelou = input_cancelavel("\nEscolha o campo para editar")
        if cancelou or opcao.lower() == "q":
            print("\nEdição cancelada.")
            time.sleep(1)
            return
        if opcao == "0":
            break

        if opcao == "1":
            v, _ = input_cancelavel(f"Novo Nome do Cliente [{dados['CLIENTE'] or ''}]")
            if v: 
                dados['CLIENTE'] = v.upper().strip()
                houve_alteracao = True
        elif opcao == "2":
            v, _ = input_cancelavel(f"Nova Parte Contrária [{dados['PARTE_CONTRARIA'] or ''}]")
            if v: 
                dados['PARTE_CONTRARIA'] = v.upper().strip()
                houve_alteracao = True
        elif opcao == "3":
            v, _ = input_cancelavel(f"Novo Cartório/Vara [{dados['CARTORIO'] or ''}]")
            if v: 
                dados['CARTORIO'] = v.upper().strip()
                houve_alteracao = True
        elif opcao == "4":
            v, _ = input_cancelavel(f"Nova Situação ([1] ATIVO / [2] CONCLUIDO) [{dados['SITUACAO']}]")
            if v == "2": 
                dados['SITUACAO'] = "CONCLUIDO"
                houve_alteracao = True
            elif v == "1": 
                dados['SITUACAO'] = "ATIVO"
                houve_alteracao = True
        elif opcao == "5":
            v, _ = input_cancelavel(f"Nova Data Distribuição (DDMMAAAA) [{dados['DISTRIBUICAO'] or ''}]")
            if v: 
                dt = formatar_data(v)
                if dt:
                    dados['DISTRIBUICAO'] = dt
                    houve_alteracao = True
                else:
                    print(f"{VERMELHO}[!] Data inválida.{RESET}")
                    time.sleep(1)
        elif opcao == "6":
            v, _ = input_cancelavel(f"Nova Observação [{dados['OBS'] or ''}]")
            if v: 
                dados['OBS'] = v.upper().strip()
                houve_alteracao = True
        else:
            print(f"{AMARELO}Opção inválida.{RESET}")
            time.sleep(1)

    if houve_alteracao:
        try:
            cursor.execute("""
                UPDATE processos SET 
                    CLIENTE = ?, PARTE_CONTRARIA = ?, CARTORIO = ?, 
                    SITUACAO = ?, DISTRIBUICAO = ?, OBS = ?, ULTIMA_VERIFICACAO = ? 
                WHERE PROCESSO = ?
            """, (
                dados['CLIENTE'], dados['PARTE_CONTRARIA'], dados['CARTORIO'], 
                dados['SITUACAO'], dados['DISTRIBUICAO'], (dados['OBS'] or "").strip(), 
                hoje(), processo_id
            ))
            conn.commit()
            print(f"\n{VERDE}[OK] Processo atualizado com sucesso!{RESET}")
        except Exception as e:
            print(f"{VERMELHO}[ERRO] Falha ao atualizar: {e}{RESET}")
    else:
        print("\nNenhuma alteração foi feita.")
    
    time.sleep(1.5)

def ver_andamento(conn: sqlite3.Connection) -> None:
    pagina = 0
    limite = 5

    while True:
        limpar_terminal()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT COUNT(*) FROM processos 
            WHERE SITUACAO = 'ATIVO' AND ULTIMA_VERIFICACAO <= date('now', '-25 days')
        """)
        total = cursor.fetchone()[0]

        if total == 0:
            print(f"\n{VERDE}[OK] Nenhum processo parado há mais de 25 dias!{RESET}")
            time.sleep(1.5)
            return

        total_paginas = max(1, (total + limite - 1) // limite)
        
        if pagina >= total_paginas:
            pagina = total_paginas - 1

        cursor.execute("""
            SELECT * FROM processos 
            WHERE SITUACAO = 'ATIVO' AND ULTIMA_VERIFICACAO <= date('now', '-25 days')
            ORDER BY ULTIMA_VERIFICACAO ASC LIMIT ? OFFSET ?
        """, (limite, pagina * limite))
        rows = cursor.fetchall()

        print(f"\n{AMARELO}-------------- PROCESSOS PARADOS (HÁ 25+ DIAS) | Pág: {pagina + 1}/{total_paginas} | TOTAL: {total} --------------{RESET}")
        for i, r in enumerate(rows):
            cartorio = r['CARTORIO'] if r['CARTORIO'] else 'N/A'
            distribuicao = r['DISTRIBUICAO'] if r['DISTRIBUICAO'] else 'N/A'
            ult_verif = r['ULTIMA_VERIFICACAO'] if r['ULTIMA_VERIFICACAO'] else 'N/A'
            cpf_cnpj = r['CPF_CNPJ'] if r['CPF_CNPJ'] else 'N/A'
            contraria = r['PARTE_CONTRARIA'] if r['PARTE_CONTRARIA'] else 'N/A'
            obs = r['OBS'] if r['OBS'] else ''

            print("-" * 80)
            print(f"{VERDE}{i + 1}.{RESET} 🏛️  CARTÓRIO: {cartorio}")
            print(f"     📄 PROC: {r['PROCESSO']} | SITUAÇÃO: {VERDE}{r['SITUACAO']}{RESET}")
            print(f"     ⚠️  ÚLT. VERIF: {ult_verif} | 📅 DISTR: {distribuicao}")
            print(f"     👤 CLIENTE: {r['CLIENTE']} | CPF/CNPJ: {cpf_cnpj}")
            print(f"     ⚔️  CONTRÁRIO: {contraria}")
            if obs:
                print(f"     📝 OBS: {obs}")
        print("-" * 80)

        print("[<] Ant | [>] Próx | [A+Nº] Atualizar Visto | [E+Nº] Editar | [G] Gerar Doc | [Q] Voltar")
        acao = input("Comando: ").strip().lower()

        if acao == "q":
            return
        elif acao in (">", ".") and pagina + 1 < total_paginas:
            pagina += 1
        elif acao in ("<", ",") and pagina > 0:
            pagina -= 1
        elif acao == "g":
            preencher_docx(conn)
        elif acao.startswith("a"):
            try:
                idx = int(acao.replace("a", "").strip()) - 1
                if 0 <= idx < len(rows):
                    cursor.execute(
                        "UPDATE processos SET ULTIMA_VERIFICACAO = ? WHERE PROCESSO = ?",
                        (hoje(), rows[idx]['PROCESSO'])
                    )
                    conn.commit()
                    print(f"{VERDE}[OK] Data de verificação atualizada!{RESET}")
                    time.sleep(0.8)
                else:
                    print(f"{VERMELHO}[!] Número inválido.{RESET}")
                    time.sleep(1)
            except ValueError:
                print(f"{VERMELHO}[!] Use 'a' + número.{RESET}")
                time.sleep(1)
        elif acao.startswith("e"):
            try:
                idx = int(acao.replace("e", "").strip()) - 1
                if 0 <= idx < len(rows):
                    editar_processo(conn, rows[idx]['PROCESSO'])
                else:
                    print(f"{VERMELHO}[!] Número inválido.{RESET}")
                    time.sleep(1)
            except ValueError:
                print(f"{VERMELHO}[!] Use 'e' + número.{RESET}")
                time.sleep(1)
        else:
            print(f"{VERMELHO}🚨 Opção inválida.{RESET}")
            time.sleep(0.8)